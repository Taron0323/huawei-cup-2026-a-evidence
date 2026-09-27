"""Extract source-located requirements and inspect every official A graph."""

import argparse
import collections
import json
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT / "src/a_solver"))
import a_solver
from evaluation_validation import validate_graph
from run_batch import sha, save_json, write_csv
from verify import dependencies


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--review-root", type=Path, required=True)
    args = parser.parse_args()
    output = args.review_root.resolve()
    model = output / "02_model"
    model.mkdir(parents=True, exist_ok=True)
    question_root = ROOT.parent / "第二十三届中国研究生数学建模竞赛 - 中文题目/中文题目"
    requirements = []
    for letter in "ABCDEF":
        source = next(p for p in sorted((question_root / f"{letter}题").glob("*.docx")) if "(1)" not in p.name)
        doc = Document(source)
        records, current = [], ""
        for i, p in enumerate(doc.paragraphs, 1):
            text = "".join(node.text or "" for node in p._p.iter()
                           if node.tag in {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t",
                                           "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"}).strip()
            if not text:
                continue
            locator = f"word/document.xml/body/p[{i}]"
            records.append({"locator": locator, "text": text, "has_math": bool(p._p.xpath(".//m:oMath"))})
            marker = re.match(r"^问题\s*([一二三四五六1-6])\s*[：:]", text)
            if marker:
                number = marker.group(1)
                current = str("一二三四五六".index(number) + 1) if number in "一二三四五六" else number
            if letter == "C":
                if text.startswith("视觉刺激产生P300"):
                    current = "1"
                elif text.startswith("建立脑电图计算模型"):
                    current = "2"
                elif text.startswith("建立认知模型"):
                    current = "3"
            if re.match(r"^(附录|算法优化方向|名词解释|参考文献|附件说明|三、)", text):
                current = ""
            if current and (marker or any(word in text for word in ("须", "请", "给出", "建立", "设计", "结果", "附件：", "正文：", "附录：", "确定", "要求"))):
                requirements.append({"requirement_id": f"{letter}-Q{current}-P{i:03d}", "problem": letter,
                    "subquestion": current, "quote": text, "source_file": str(source.relative_to(ROOT.parent)),
                    "source_page_or_locator": locator, "expected_output": "逐段回应题面并保留指定结果格式",
                    "data_dependency": f"{letter}题随题附件", "verification_plan": "逐段对账与结果文件定位",
                    "status": "PLANNED" if letter == "A" else "SELECTION_REVIEW_ONLY",
                    "notes": "正文与OMML文字按原顺序保存；复杂公式以原始DOCX为准"})
        save_json(model / f"{letter}_source_paragraphs.json", {"source": str(source.relative_to(ROOT.parent)),
                   "sha256": sha(source), "paragraphs": records,
                   "tables": [[[[c.text for c in row.cells] for row in table.rows]] for table in doc.tables]})
    write_csv(model / "all_requirements.csv", requirements)
    selected = [row for row in requirements if row["problem"] == "A"]
    write_csv(model / "requirements_trace.csv", selected)
    evidence_rows = [{"id": row["requirement_id"], "quote": row["quote"],
        "source_file": "inputs/problem/A题/通用神经网络处理器下的多核调度问题.docx",
        "source_page": row["source_page_or_locator"], "deliverable": row["expected_output"],
        "paper_location": f"问题{row['subquestion']}", "result_file": "", "status": "PLANNED",
        "review_note": row["notes"]} for row in selected]
    write_csv(ROOT / "evidence/requirements.csv", evidence_rows)

    rows = []
    for path in sorted(a_solver.DATA_ROOT.glob("case_*.json")):
        graph = a_solver.load_graph(path)
        validate_graph(graph)
        eligible, edges = dependencies(graph)
        adj = {node: set() for node in eligible}
        for a, b in edges:
            adj[a].add(b)
            adj[b].add(a)
        seen, sizes = set(), []
        for node in sorted(eligible):
            if node in seen:
                continue
            seen.add(node)
            stack, count = [node], 0
            while stack:
                source = stack.pop()
                count += 1
                nxt = adj[source] - seen
                seen.update(nxt)
                stack.extend(nxt)
            sizes.append(count)
        by_pipe = collections.Counter()
        for op in graph["ops"]:
            if op["id"] in eligible:
                by_pipe[op["pipe"]] += max(1, op["cycles"])
        rows.append({"case": path.stem, "sha256": sha(path), "bytes": path.stat().st_size,
                     "operations": len(graph["ops"]), "eligible_operations": len(eligible),
                     "tensors": len(graph["tensors"]), "edges": len(graph["edges"]),
                     "components": len(sizes), "largest_component": max(sizes, default=0),
                     "compute_cycles": sum(by_pipe.values()), "cube_cycles": by_pipe["PIPE_M"],
                     "vector_cycles": by_pipe["PIPE_V"],
                     "zero_size_tensors": sum(t["size"] == 0 for t in graph["tensors"]),
                     "zero_cycle_operations": sum(o["cycles"] == 0 for o in graph["ops"]),
                     "missing_required_fields": 0, "duplicate_node_ids": 0, "duplicate_edges": 0,
                     "invalid_endpoints": 0, "cycles_in_graph": 0, "status": "AI_VERIFIED"})
    write_csv(model / "data_health.csv", rows)
    report = ("# A题数据体检\n\n" +
        f"检查{len(rows)}个原始计算图，逐图运行附件的格式、节点ID、边端点、无环、非负数值和流水线字段验证。\n\n" +
        f"核内操作共{sum(r['eligible_operations'] for r in rows)}个；张量共{sum(r['tensors'] for r in rows)}个。\n\n" +
        "所有图均可读取并通过字段检查。零大小张量与零周期操作保留，执行时依照官方评估器的最小一周期规则处理。未进行删除、插值、缩放或单位转换。\n\n" +
        "size的单位为bytes，cycles和Makespan的单位为clock cycles；数据不包含实测GHz，不能换算为秒。\n\n" +
        "不同图的弱连通分量数量与最大分量差异很大，完整统计见data_health.csv。逐图机器验证不构成现实硬件测量。\n")
    (model / "data_health_report.md").write_text(report, encoding="utf-8")
    print(json.dumps({"graphs": len(rows), "requirements_A": len(selected), "requirements_all": len(requirements)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
