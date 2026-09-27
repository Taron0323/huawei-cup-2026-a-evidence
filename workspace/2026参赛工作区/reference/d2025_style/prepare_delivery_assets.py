"""Prepare source snapshots and an official-layout abstract for the new template."""
import datetime
import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
WORKSPACE = HERE.parents[1]
TARGET = WORKSPACE / "paper/latex_d2025_style"


def main():
    sources = TARGET / "sources"
    sources.mkdir(exist_ok=True)
    shutil.copyfile(HERE / "reference.pdf", sources / "reference.pdf")
    shutil.copyfile(HERE / "layout-evidence.json", sources / "layout-evidence.json")
    provenance = []
    for source, name, url in [
        (Path("/tmp/huawei-d2025-web.html"), "problem-cnblogs.html", "https://www.cnblogs.com/li-zi-ke-yan/articles/19103763"),
        (Path("/tmp/huawei-bing2.html"), "search-bing.html", "https://cn.bing.com/search?q=%22%E4%BD%8E%E7%A9%BA%E6%B9%8D%E6%B5%81%22&count=10"),
    ]:
        data = source.read_bytes()
        (sources / name).write_bytes(data)
        provenance.append({"file": name, "url": url, "retrieved_on": "2026-09-21", "archived_at": datetime.datetime.now().astimezone().isoformat(),
                           "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data),
                           "source_type": "public_reprint" if "problem" in name else "search_result"})
    soup = BeautifulSoup((sources / "problem-cnblogs.html").read_text(), "html.parser")
    text = soup.select_one("#cnblogs_post_body").get_text("\n", strip=True)
    start, end = text.index("低空湍流监测及最优航路规划"), text.index("📚2 完整资源下载")
    (sources / "2025D-题面转载摘录.txt").write_text(text[start:end], encoding="utf-8")
    (sources / "provenance.json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    rules = sources / "2026-rules"
    rules.mkdir(exist_ok=True)
    attachments = WORKSPACE / "official/20260921/02_官方附件"
    for pattern in ["附件2*", "附件4*"]:
        for file in attachments.glob(pattern):
            shutil.copyfile(file, rules / file.name)
    shutil.copyfile(WORKSPACE / "paper/2026格式与合规对照.md", rules / "2026格式与合规对照.md")


def prepare_frontmatter():
    doc = Document(WORKSPACE / "paper/word/working-template.docx")
    for p in list(doc.paragraphs)[24:]:
        p._element.getparent().remove(p._element)
    doc.paragraphs[17].text = "题 目：基于多源数据融合的低空湍流监测与航路优化"
    abstract = [
        "低空飞行所处的大气环境具有局地变化快、垂直结构复杂和观测尺度不一致的特点。围绕湍流强度评估、三维场重构与航路优化三个相互衔接的问题，本文以多源观测为基础，将物理诊断、统计标定和图搜索相结合，构建从观测参数到飞行风险的计算框架。",
        "问题一：针对风廓线雷达与微波辐射计在观测要素上的互补关系，首先完成时间匹配、高度对齐与质量控制，计算位温、风切变和理查逊数。融合热力稳定性与雷达谱宽代理量建立模型 a，再仅使用雷达变量及历史变化构建模型 b。以模型 a 为参照，采用约束回归和随机森林进行标定，并按时间块比较垂直廓线和误差分布。测试集 RMSE、MAE 与决定系数分别为【待计算】、【待计算】和【待计算】，相对未标定模型的误差变化为【待计算】。该比较描述模型 b 对参照指标的逼近程度，指标的物理有效性另用独立湍流观测检验。",
        "在参数准备阶段，以位温梯度刻画热力层结，以水平风分量的垂直梯度刻画动力切变，由二者构造理查逊数。雷达谱宽先完成可获得的非湍流展宽校正，再作为湍流活动的代理变量参与综合评价。稳定度映射采用有界函数，归一化参数在训练阶段确定，并在验证阶段保持固定。通过分高度残差比较和权重扰动，分析不同变量对综合指标的作用及其不确定性。",
        "问题二：将地面自动站、风廓线雷达及天气雷达转换到统一时间和坐标系统，以各向异性距离、设备误差和时间衰减构造融合权重。进一步建立包含观测项、背景项和平滑项的变分模型，重构高度低于 2 km、水平间距 100 m、垂直间距 50 m 的三维湍流场。通过留站验证、设备消融及分高度统计，得到【主要空间特征】与【可靠范围】，重构 RMSE 为【待计算】，区间覆盖率为【待计算】。",
        "为使不同设备在各自的代表尺度上提供约束，通过观测算子将目标网格映射到点位、垂直门或雷达采样体积。线性观测关系与二次目标形成稀疏线性系统，按其正定条件选择共轭梯度等求解方法；存在指标范围约束时采用有界优化。时间更新结合状态估计与半拉格朗日平流，并使用分块重采样描述误差范围。最终同时展示风险分布、有效观测覆盖和不确定性，区分观测充分区域与主要依赖背景信息的区域。",
        "问题三：在有数值预报和无数值预报两类条件下，分别构建模型 d 与模型 e。模型 d 利用 02:00 至 05:00 的重构场标定 NWP 诊断量，生成 05:00 至 08:00 的预报；模型 e 利用截至 05:00 的观测建立非线性外推关系，预测 05:00 至 06:00 的场。通过滚动回报检验比较各时效的误差，再将预报风险转化为时空图上的非负边代价，在给定起终点和允许区域内搜索累计湍流暴露量最小的航路。最终路径的暴露量为【待计算】，相对基准变化为【待计算】。",
        "航路层将空间位置与到达时间共同表示为图节点，以边段长度和沿线平均风险的乘积构造代价。采用 A* 搜索时，启发函数由单位距离风险下界和剩余距离共同确定；下界为零时采用 Dijkstra 搜索。对所得路径重新计算完整边段的可行性和风险积分，并通过网格加密、参数扰动和预报情景变化，评价路径选择对离散尺度及场误差的敏感程度。路径最优性在给定离散图、允许区域和风险定义内讨论。",
        "三个问题分别形成单站廓线、区域风险场与预报航路，各阶段输出对应后续输入和验证参照。交付内容包括逐点结果、三维网格、路径节点序列及可运行程序，所有数值按实际计算结果填写。",
    ]
    blanks = [doc.paragraphs[i] for i in [20, 21, 22]]
    anchor = blanks[0]
    for text in abstract:
        p = anchor.insert_paragraph_before(text)
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(5)
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        p._element.get_or_add_pPr().append(snap)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    for blank in blanks:
        blank._element.getparent().remove(blank._element)
    for p in doc.paragraphs:
        if p.text.startswith("题 目："):
            for run in p.runs:
                run.font.size = Pt(16)
                run.font.name = "SimHei"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        if p.text.startswith("关键词："):
            p.text = "关键词：湍流强度量化；多源数据融合；三维场重构；航路规划；变分同化"
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.name = "Times New Roman"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            p.paragraph_format.line_spacing = 1.0
            snap = OxmlElement("w:snapToGrid")
            snap.set(qn("w:val"), "0")
            p._element.get_or_add_pPr().append(snap)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(TARGET / "frontmatter.docx")
    env = dict(os.environ, FONTCONFIG_FILE=str(WORKSPACE / "tools/fonts.conf"))
    subprocess.run([shutil.which("soffice"), "-env:UserInstallation=file:///tmp/huawei-d2025-lo",
                    "--headless", "--convert-to", "pdf", "--outdir", str(TARGET), str(TARGET / "frontmatter.docx")],
                   check=True, env=env)
    print("Official front matter pages:", len(PdfReader(TARGET / "frontmatter.pdf").pages))


if __name__ == "__main__":
    if "--frontmatter-only" not in sys.argv:
        main()
    prepare_frontmatter()
