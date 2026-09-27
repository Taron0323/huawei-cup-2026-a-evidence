#!/usr/bin/env python3
"""Extract question text and lightweight metadata without unpacking large archives."""

from __future__ import annotations

import argparse
import csv
import json
import lzma
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document


def canonical_doc(problem_dir: Path) -> Path | None:
    docs = sorted(p for p in problem_dir.glob("*.docx") if "(1)" not in p.name)
    return docs[0] if docs else None


def extract_docx(path: Path, output: Path) -> dict[str, object]:
    document = Document(path)
    lines: list[str] = [f"# {path.stem}", "", f"Source: `{path}`", ""]
    paragraphs = 0
    question_markers: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraphs += 1
        lines.append(text)
        if re.match(r"^(问题|附录|[一二三四五六七八九十]+[、.])", text):
            question_markers.append(text)
    for index, table in enumerate(document.tables, 1):
        lines.extend(["", f"## Table {index}", ""])
        for row in table.rows:
            lines.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {
        "source": str(path),
        "paragraphs": paragraphs,
        "tables": len(document.tables),
        "question_markers": question_markers,
    }


def zip_summary(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        infos = archive.infolist()
        files = [item for item in infos if not item.is_dir()]
        extensions = Counter(Path(item.filename).suffix.lower() or "[no_ext]" for item in files)
        return {
            "path": str(path),
            "members": len(infos),
            "files": len(files),
            "uncompressed_bytes": sum(item.file_size for item in files),
            "integrity": "PASS" if bad is None else f"FAIL:{bad}",
            "extensions": dict(sorted(extensions.items())),
            "sample_members": [item.filename for item in infos[:20]],
        }


def xlsx_summary(path: Path) -> dict[str, object]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        rows = sheet.iter_rows(min_row=1, max_row=5, values_only=True)
        sample = [[str(value)[:160] if value is not None else "" for value in row] for row in rows]
        sheets.append({"name": name, "rows": sheet.max_row, "columns": sheet.max_column, "sample": sample})
    return {"path": str(path), "sheets": sheets}


def mat_summary(path: Path) -> dict[str, object]:
    from scipy.io import loadmat

    try:
        values = loadmat(path, squeeze_me=False, struct_as_record=False)
        variables = []
        for name, value in values.items():
            if name.startswith("__"):
                continue
            variables.append({"name": name, "shape": list(getattr(value, "shape", ())), "dtype": str(getattr(value, "dtype", ""))})
        return {"path": str(path), "format": "MAT", "variables": variables[:80]}
    except Exception as exc:
        return {"path": str(path), "format": "MAT", "status": "UNVERIFIED", "error": f"{type(exc).__name__}: {exc}"}


def csv_summary(path: Path) -> dict[str, object]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        rows = []
        for row in reader:
            rows.append(row[:20])
            if len(rows) >= 4:
                break
    return {"path": str(path), "sample_rows": rows}


def json_summary(path: Path) -> dict[str, object]:
    try:
        with path.open("r", encoding="utf-8", errors="replace") as handle:
            value = json.load(handle)
        if isinstance(value, dict):
            return {"path": str(path), "kind": "object", "keys": sorted(value)[:80]}
        return {"path": str(path), "kind": type(value).__name__, "length": len(value) if hasattr(value, "__len__") else None}
    except Exception as exc:
        return {"path": str(path), "status": "UNVERIFIED", "error": f"{type(exc).__name__}: {exc}"}


def jsonl_xz_summary(path: Path) -> dict[str, object]:
    sample: list[str] = []
    lines = 0
    with lzma.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            lines += 1
            if len(sample) < 3:
                sample.append(line[:500].rstrip())
            if lines >= 1000:
                break
    return {"path": str(path), "sampled_lines": lines, "first_lines": sample, "full_count": "NOT_COUNTED"}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    root = args.root.expanduser().resolve()
    output = args.output.expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    summaries: dict[str, object] = {}
    for letter in "ABCDEF":
        problem_dir = root / f"{letter}题"
        problem_output = output / f"{letter}_question.md"
        doc = canonical_doc(problem_dir)
        summaries[letter] = {"directory": str(problem_dir), "bytes": sum(p.stat().st_size for p in problem_dir.rglob("*") if p.is_file())}
        if doc:
            summaries[letter]["question"] = extract_docx(doc, problem_output)
        else:
            (output / f"{letter}_question.md").write_text(f"# {letter}题\n\nNo canonical DOCX found.\n", encoding="utf-8")
        files = [p for p in problem_dir.rglob("*") if p.is_file() and "(1)" not in p.name and p.name != ".DS_Store"]
        metadata: list[dict[str, object]] = []
        for path in sorted(files):
            suffix = path.suffix.lower()
            if suffix == ".zip":
                metadata.append(zip_summary(path))
            elif suffix == ".xlsx":
                metadata.append(xlsx_summary(path))
            elif suffix == ".mat":
                metadata.append(mat_summary(path))
            elif suffix == ".csv":
                metadata.append(csv_summary(path))
            elif suffix == ".json" and path.stat().st_size <= 2_000_000:
                metadata.append(json_summary(path))
            elif suffix == ".xz":
                metadata.append(jsonl_xz_summary(path))
        (output / f"{letter}_metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        summaries[letter]["metadata_records"] = len(metadata)
    (output / "question_summary.json").write_text(json.dumps(summaries, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summaries, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
